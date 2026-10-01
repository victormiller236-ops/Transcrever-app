#!/usr/bin/env python3
"""Cruza os IDs da planilha de trilhos com a planilha de IDs de bobinas (SAP).

Uso:
    pip install openpyxl
    python cruzar_ids.py TRILHOS.xlsx SAP.xlsx [-o saida.xlsx]

Gera um Excel com as abas: Trilhos (cópia), Cruzamento (IDs que coincidiram)
e Nao encontrados (IDs dos trilhos que não existem no SAP).
"""
import argparse
import datetime as dt
import sys
import unicodedata

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# (título na saída, origem "T"=trilhos / "S"=SAP, coluna de origem normalizada)
SAIDA = [
    ("Trilho", "T", "trilho"),
    ("ID", "S", "id de material"),
    ("Data de criação", "S", "criado em"),
    ("Material", "S", "material"),
    ("Texto breve de material", "S", "texto breve de material"),
    ("Largura (mm)", "S", "largura (mm)"),
    ("Peso Atual", "S", "peso atual"),
    ("Peso Original", "S", "peso original"),
    ("Designação da Chapa", "S", "designacao da chapa"),
    ("Forma de Criação", "S", "forma de criacao"),
    ("Nome completo", "S", "nome completo"),
    ("Bloqueado", "S", "bloqueado"),
    ("Status do ID de Material", "S", "status do id de material"),
    ("ID da bobina de origem", "B", "id"),
    ("Bobina de origem (texto breve)", "B", "texto"),
]
LARGURAS = {"Bobina de origem (texto breve)": 42, "Texto breve de material": 42, "Nome completo": 30}


def norm(s):
    s = unicodedata.normalize("NFD", str(s if s is not None else ""))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return " ".join(s.strip().lower().split())


def norm_id(v):
    if v is None:
        return ""
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s.lstrip("0") or ("0" if s else "")


def ler_aba(caminho, coluna_chave):
    """Devolve (cabeçalho, linhas) da primeira aba que tenha a coluna-chave."""
    wb = load_workbook(caminho, read_only=True, data_only=True)
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        cab = next(it, None)
        if cab and coluna_chave in [norm(c) for c in cab]:
            return list(cab), [r for r in it if any(c is not None for c in r)]
    sys.exit(f'Erro: coluna "{coluna_chave}" não encontrada em {caminho}')


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("trilhos")
    ap.add_argument("sap")
    ap.add_argument("-o", "--saida", default=f"cruzamento-ids-{dt.date.today().isoformat()}.xlsx")
    a = ap.parse_args()

    print("Lendo trilhos...")
    cab_t, rows_t = ler_aba(a.trilhos, "id")
    print("Lendo SAP (pode levar um pouco)...")
    cab_s, rows_s = ler_aba(a.sap, "id de material")
    ct = {norm(c): i for i, c in reversed(list(enumerate(cab_t)))}
    cs = {norm(c): i for i, c in reversed(list(enumerate(cab_s)))}
    if "trilho" not in ct:
        sys.exit('Erro: a planilha de trilhos precisa ter a coluna "Trilho".')
    faltando = [t for t, o, k in SAIDA if o == "S" and k not in cs] + ([] if "id de material original" in cs else ["ID da bobina de origem"])
    if faltando:
        print("Aviso: colunas não encontradas no SAP (ficarão vazias):", ", ".join(faltando))

    idx = {}
    for r in rows_s:
        i = norm_id(r[cs["id de material"]])
        if i and i not in idx:
            idx[i] = r

    out, nao, vistos, por_id = [], [], set(), {}
    for r in rows_t:
        i = norm_id(r[ct["id"]])
        if not i:
            continue
        trilho = r[ct["trilho"]]
        if (i, trilho) in vistos:
            continue
        vistos.add((i, trilho))
        s = idx.get(i)
        if s is None:
            nao.append((trilho, i))
            continue
        linha = []
        for _, origem, k in SAIDA:
            if origem == "T":
                linha.append(trilho)
            elif origem == "B":
                o = norm_id(s[cs["id de material original"]]) if "id de material original" in cs else ""
                if not o or o == "0":
                    linha.append("")
                elif k == "id":
                    linha.append(o)
                else:
                    bo = idx.get(o)
                    linha.append(bo[cs["texto breve de material"]] if bo else "")
            elif k == "id de material":
                linha.append(i)
            else:
                linha.append(s[cs[k]] if k in cs else None)
        out.append(linha)
        por_id.setdefault(i, []).append(trilho)

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Trilhos"
    ws0.append(cab_t)
    for r in rows_t:
        ws0.append(list(r))

    ws = wb.create_sheet("Cruzamento")
    ws.append([t for t, _, _ in SAIDA])
    for l in out:
        ws.append(l)
    cab_fill = PatternFill("solid", fgColor="0B6E4F")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = cab_fill
        c.alignment = Alignment(vertical="center")
    for n, (t, _, _) in enumerate(SAIDA, 1):
        ws.column_dimensions[get_column_letter(n)].width = LARGURAS.get(t, max(11, len(t) + 2))
        for c in ws[get_column_letter(n)][1:]:
            if isinstance(c.value, (dt.datetime, dt.date)):
                c.number_format = "dd/mm/yyyy"
            elif t in ("Largura (mm)", "Peso Atual", "Peso Original"):
                c.number_format = "#,##0.###"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    if nao:
        w3 = wb.create_sheet("Nao encontrados")
        w3.append(["Trilho", "ID"])
        for l in nao:
            w3.append(list(l))

    wb.save(a.saida)

    print(f"\nIDs lidos nos trilhos : {len(vistos)}")
    print(f"Coincidiram com o SAP : {len(out)}")
    print(f"Não encontrados no SAP: {len(nao)}" + "".join(f"\n   ID {i} (trilho {t})" for t, i in nao))
    rep = {i: t for i, t in por_id.items() if len(t) > 1}
    if rep:
        print("IDs em mais de um trilho:", ", ".join(f"{i} (trilhos {', '.join(map(str, t))})" for i, t in rep.items()))
    print(f"\nArquivo gerado: {a.saida}")


if __name__ == "__main__":
    main()
